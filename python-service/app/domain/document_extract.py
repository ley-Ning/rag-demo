"""多格式文档文本提取：PDF / Word(docx) / HTML / 纯文本。

- PDF：pypdf 逐页提取文本（仅文本型 PDF；扫描件需要 OCR，明确报错提示）
- docx：python-docx 按段落顺序提取，表格按行拼接为制表符分隔文本
- HTML：复用 web 插件同款正则清洗（去 script/style/标签 + 实体解码）
- 文本类：多编码尝试（utf-8 / utf-8-sig / gb18030）

统一由 SUPPORTED_EXTRACT_EXTENSIONS 声明支持范围，worker 与上传端点共用。
"""

import logging
from pathlib import Path

logger = logging.getLogger(__name__)

TEXT_EXTENSIONS = {".txt", ".md", ".markdown", ".text", ".log", ".csv", ".json"}
PDF_EXTENSIONS = {".pdf"}
DOCX_EXTENSIONS = {".docx"}
HTML_EXTENSIONS = {".html", ".htm"}

SUPPORTED_EXTRACT_EXTENSIONS = TEXT_EXTENSIONS | PDF_EXTENSIONS | DOCX_EXTENSIONS | HTML_EXTENSIONS


def _extract_text_bytes(raw: bytes) -> str:
    for encoding in ("utf-8", "utf-8-sig", "gb18030"):
        try:
            return raw.decode(encoding)
        except Exception:
            continue
    raise RuntimeError("文件解码失败，请确保是 UTF-8 或 GB18030 文本文件")


def _extract_pdf(file_path: Path) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise RuntimeError("PDF 解析依赖未安装（pip install pypdf）") from exc

    reader = PdfReader(str(file_path))
    pages: list[str] = []
    for index, page in enumerate(reader.pages, start=1):
        try:
            text = (page.extract_text() or "").strip()
        except Exception as exc:
            logger.warning("PDF page %s extract failed: %s", index, exc)
            text = ""
        if text:
            pages.append(f"[第{index}页]\n{text}")
    if not pages:
        raise RuntimeError(
            "PDF 未提取到文本：可能是扫描件/图片型 PDF，请先做 OCR 或改传文本版"
        )
    return "\n\n".join(pages)


def _extract_docx(file_path: Path) -> str:
    try:
        from docx import Document
    except ImportError as exc:
        raise RuntimeError("Word 解析依赖未安装（pip install python-docx）") from exc

    document = Document(str(file_path))
    parts: list[str] = [paragraph.text.strip() for paragraph in document.paragraphs]
    parts = [part for part in parts if part]

    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                parts.append("\t".join(cells))

    if not parts:
        raise RuntimeError("Word 文档未提取到文本内容")
    return "\n".join(parts)


def _extract_html(raw: bytes) -> str:
    import html as html_lib
    import re

    content = raw.decode("utf-8", errors="ignore")
    content = re.sub(r"(?is)<script[^>]*>.*?</script>", " ", content)
    content = re.sub(r"(?is)<style[^>]*>.*?</style>", " ", content)
    content = re.sub(r"(?is)<noscript[^>]*>.*?</noscript>", " ", content)
    content = re.sub(r"(?is)<[^>]+>", " ", content)
    content = html_lib.unescape(content)
    content = re.sub(r"[ \t]+", " ", content)
    content = re.sub(r"\n\s*\n+", "\n\n", content)
    text = content.strip()
    if not text:
        raise RuntimeError("HTML 文件未提取到文本内容")
    return text


def extract_text(file_path: Path, extension: str) -> str:
    """按扩展名分发提取文本；不支持的类型抛 RuntimeError（消息面向用户）"""
    ext = extension.lower()
    raw: bytes | None = None

    if ext in TEXT_EXTENSIONS:
        raw = file_path.read_bytes()
        return _extract_text_bytes(raw)
    if ext in PDF_EXTENSIONS:
        return _extract_pdf(file_path)
    if ext in DOCX_EXTENSIONS:
        return _extract_docx(file_path)
    if ext in HTML_EXTENSIONS:
        raw = file_path.read_bytes()
        return _extract_html(raw)
    raise RuntimeError(
        f"暂不支持的文件类型: {ext or 'unknown'}，"
        f"当前支持 {'/'.join(sorted(SUPPORTED_EXTRACT_EXTENSIONS))}"
    )
